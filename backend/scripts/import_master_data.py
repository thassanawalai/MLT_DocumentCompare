"""One-time import of the current HS Code CSV and Freetime workbook."""
import csv
from datetime import date, datetime
import os
from pathlib import Path
import re
import sys
import zipfile
import argparse

from sqlalchemy import func, select

BACKEND_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = BACKEND_DIR.parent
sys.path.insert(0, str(BACKEND_DIR))

from app.db import SessionLocal  # noqa: E402
from app.models.db_models import FreetimeRecord, HSCodeRecord  # noqa: E402


DATA_DIR = (
    Path(os.environ["MASTER_DATA_DIR"])
    if "MASTER_DATA_DIR" in os.environ
    else REPO_ROOT / "frontend" / "public"
)
HS_CSV = DATA_DIR / "HS_Master.csv"
FREETIME_XLSX = DATA_DIR / "freetime_data.xlsx"


def text_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_hs_records() -> list[dict[str, str]]:
    with HS_CSV.open("r", encoding="utf-8-sig", newline="") as source:
        rows = csv.reader(source)
        try:
            headers = [header.strip().casefold() for header in next(rows)]
        except StopIteration as exc:
            raise ValueError(f"HS Code source is empty: {HS_CSV}") from exc

        required = {"sales", "customer name", "commodity", "hs code"}
        missing = required.difference(headers)
        if missing:
            raise ValueError(f"HS Code CSV is missing columns: {', '.join(sorted(missing))}")
        positions = {name: headers.index(name) for name in required}

        records = []
        for row_number, row in enumerate(rows, start=2):
            if not any(cell.strip() for cell in row):
                continue

            def cell(header: str) -> str:
                index = positions[header]
                return row[index].strip() if index < len(row) else ""

            records.append({
                "sale": cell("sales"),
                "customer": cell("customer name"),
                "commodity": cell("commodity"),
                "hs_code": cell("hs code"),
            })
    return records


def normalize_header(value) -> str:
    return re.sub(r"\s+", " ", text_value(value).strip()).casefold()


def read_freetime_records() -> list[dict[str, str]]:
    workbook = None
    source = None
    if zipfile.is_zipfile(FREETIME_XLSX):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to import the XLSX Freetime source") from exc
        workbook = load_workbook(FREETIME_XLSX, read_only=True, data_only=True)
        rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
    else:
        # The current .xlsx-named source is actually comma-separated text.
        source = FREETIME_XLSX.open("r", encoding="utf-8-sig", newline="")
        rows = csv.reader(source)

    try:
        try:
            headers = [normalize_header(value) for value in next(rows)]
        except StopIteration as exc:
            raise ValueError(f"Freetime source is empty: {FREETIME_XLSX}") from exc

        source_columns = {
            "carrier": ("carrier",),
            "lat_krabang": ("lat krabang",),
            "laem_chabang": ("laem chabang",),
            "pat_bangkok": ("pat (bangkok)",),
            "songkla": ("songkla",),
            "reefer": ("reefer",),
            "remark": ("remark",),
            "free_time_plug": ("free time plug",),
            "dem_20_40": ("dem 20'gp & 40'hc",),
            "dt_20_40": ("dt 20'gp & 40'hc",),
            "dem_reefer": ("dem/reffer",),
            "det_reefer": ("det/reffeer",),
        }
        positions = {}
        for field, aliases in source_columns.items():
            found = next((headers.index(alias) for alias in aliases if alias in headers), None)
            if found is not None:
                positions[field] = found
        required = set(source_columns)
        missing = required.difference(positions)
        if missing:
            raise ValueError(f"Freetime workbook is missing columns: {', '.join(sorted(missing))}")

        records = []
        last_carrier = ""
        for row in rows:
            values = {field: text_value(row[index] if index < len(row) else "")
                      for field, index in positions.items()}
            if not any(values.values()):
                continue
            if values["carrier"]:
                last_carrier = values["carrier"]
            elif last_carrier:
                # The current workbook uses blank carrier cells for continued carrier groups.
                values["carrier"] = last_carrier

            records.append(values)
        return records
    finally:
        if workbook is not None:
            workbook.close()
        if source is not None:
            source.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--if-empty",
        action="store_true",
        help="skip if both tables already match the source row counts",
    )
    args = parser.parse_args()
    if not HS_CSV.is_file():
        raise FileNotFoundError(f"HS Code source not found: {HS_CSV}")
    if not FREETIME_XLSX.is_file():
        raise FileNotFoundError(f"Freetime source not found: {FREETIME_XLSX}")

    hs_records = read_hs_records()
    freetime_records = read_freetime_records()
    if not hs_records or not freetime_records:
        raise ValueError("Import stopped because one or both source files contain no records")

    with SessionLocal.begin() as db:
        hs_count = db.scalar(select(func.count()).select_from(HSCodeRecord))
        freetime_count = db.scalar(select(func.count()).select_from(FreetimeRecord))
        if hs_count or freetime_count:
            if args.if_empty and hs_count == len(hs_records) and freetime_count == len(freetime_records):
                print(
                    "Import skipped: target tables already match the source row counts "
                    f"(hs_code_records={hs_count}, freetime_records={freetime_count})."
                )
                return
            raise RuntimeError(
                "Import stopped: target tables are not empty "
                f"(hs_code_records={hs_count}, freetime_records={freetime_count}). "
                "No rows were changed."
            )

        db.execute(HSCodeRecord.__table__.insert(), hs_records)
        db.execute(FreetimeRecord.__table__.insert(), freetime_records)

    print(f"Imported {len(hs_records)} HS Code rows into hs_code_records.")
    print(f"Imported {len(freetime_records)} Freetime rows into freetime_records.")


if __name__ == "__main__":
    main()
